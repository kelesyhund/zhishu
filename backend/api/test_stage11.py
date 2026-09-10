import tempfile
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from docx import Document as DocxDocument
from rest_framework.authtoken.models import Token
from rest_framework.test import APIClient

from api.models import Document, DocumentProcessingTask, DocumentSection, KnowledgeBase, Paragraph
from api.services.document_chunking import chunking_signature, preview_document_chunks
from api.services.document_parser import BlockType, parse_document
from api.services.document_processor import process_document
from api.services.model_resolution import document_needs_reprocess
from api.services.retrieval import retrieve_candidates
from api.services.retrieval_evaluation import EvaluationDataError
from api.services.stage11_evaluation import (
    StableEvidence,
    Stage11EvaluationCase,
    aggregate_answer_rows,
    evaluate_answer_record,
    evaluate_stable_ranking,
    load_stage11_cases,
)


class StageElevenParserTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix="stage11-parser-")
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_txt_blocks_and_ids_are_stable(self):
        path = self.root / "stable.txt"
        path.write_text("第一段。\n\n第二段。", encoding="utf-8")
        first = parse_document(str(path), "stage11-source")
        second = parse_document(str(path), "stage11-source")
        self.assertEqual([item.block_id for item in first.blocks], [item.block_id for item in second.blocks])
        self.assertEqual([item.content_sha256 for item in first.blocks], [item.content_sha256 for item in second.blocks])

    def test_markdown_preserves_heading_code_table_and_list(self):
        path = self.root / "guide.md"
        path.write_text(
            "# 部署\n\n## 网络\n\n说明。\n\n```bash\ndocker compose up\necho ok\n```\n\n"
            "| 参数 | 含义 |\n|---|---|\n| -p | 端口 |\n\n- 检查网络\n- 检查服务名\n",
            encoding="utf-8",
        )
        parsed = parse_document(str(path), "stage11-markdown")
        by_type = {item.block_type: item for item in parsed.blocks}
        self.assertEqual(by_type[BlockType.PARAGRAPH].heading_path, ("部署", "网络"))
        self.assertIn("docker compose up\necho ok", by_type[BlockType.CODE].text)
        self.assertIn("| -p | 端口 |", by_type[BlockType.TABLE].text)
        self.assertIn("- 检查服务名", by_type[BlockType.LIST].text)

    def test_docx_heading_and_table_are_structured(self):
        path = self.root / "guide.docx"
        source = DocxDocument()
        source.add_heading("部署", level=1)
        source.add_paragraph("先检查守护进程。")
        table = source.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "参数"
        table.cell(0, 1).text = "含义"
        table.cell(1, 0).text = "--restart"
        table.cell(1, 1).text = "重启策略"
        source.save(path)
        parsed = parse_document(str(path), "stage11-docx")
        paragraph = next(item for item in parsed.blocks if item.block_type == BlockType.PARAGRAPH)
        table_block = next(item for item in parsed.blocks if item.block_type == BlockType.TABLE)
        self.assertEqual(paragraph.heading_path, ("部署",))
        self.assertIn("--restart | 重启策略", table_block.text)

    @patch("api.services.document_parser.PdfReader")
    def test_pdf_preserves_page_numbers(self, reader_class):
        path = self.root / "guide.pdf"
        path.write_bytes(b"stage11-pdf-fixture")
        reader_class.return_value.pages = [
            SimpleNamespace(extract_text=lambda: "第一页内容"),
            SimpleNamespace(extract_text=lambda: "第二页内容"),
        ]
        parsed = parse_document(str(path), "stage11-pdf")
        content_blocks = [item for item in parsed.blocks if item.text]
        self.assertEqual([item.page_start for item in content_blocks], [1, 2])

    @patch("api.services.document_parser.PdfReader")
    def test_scanned_pdf_has_clear_error(self, reader_class):
        path = self.root / "scan.pdf"
        path.write_bytes(b"stage11-scan-fixture")
        reader_class.return_value.pages = [SimpleNamespace(extract_text=lambda: "")]
        with self.assertRaisesRegex(ValueError, "扫描版 PDF"):
            parse_document(str(path), "stage11-scan")


class StageElevenDocumentFlowTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix="stage11-media-")
        self.override = override_settings(MEDIA_ROOT=self.temp_dir.name)
        self.override.enable()
        self.owner = User.objects.create_user("stage11-owner", password="pass1234")
        self.other = User.objects.create_user("stage11-other", password="pass1234")
        self.knowledge = KnowledgeBase.objects.create(name="stage11-kb", owner=self.owner)
        self.document = Document.objects.create(
            knowledge_base=self.knowledge,
            name="stage11-guide.md",
            file=SimpleUploadedFile(
                "stage11-guide.md",
                b"# Network\n\nDocker services share a network.\n\n## Debug\n\nCheck the service name and network.",
            ),
            chunk_strategy=Document.ChunkStrategy.PARENT_CHILD,
            parent_max_tokens=400,
            child_target_tokens=100,
            child_overlap_tokens=20,
        )

    def tearDown(self):
        self.override.disable()
        self.temp_dir.cleanup()

    @patch("api.services.document_processor.embed_texts")
    def test_parent_sections_and_retrievable_children_are_separate(self, embed_texts):
        embed_texts.side_effect = lambda texts, knowledge: [[1.0, 0.0] for _ in texts]
        process_document(self.document)
        self.document.refresh_from_db()
        self.assertGreater(self.document.parent_chunk_count, 0)
        self.assertEqual(self.document.paragraph_count, self.document.paragraphs.count())
        self.assertEqual(self.document.parent_chunk_count, self.document.sections.count())
        self.assertTrue(self.document.paragraphs.filter(chunk_type=Paragraph.ChunkType.CHILD).exists())
        for child in self.document.paragraphs.select_related("parent_section"):
            self.assertIsNotNone(child.parent_section)
            self.assertEqual(child.parent_section.document_id, self.document.id)

    @patch("api.services.document_chunking.build_chunking_plan")
    @patch("api.services.document_chunking.parse_document")
    def test_preview_has_no_database_embedding_or_task_side_effect(self, parser, chunker):
        from api.services.document_parser import ChunkingPlan, ParsedDocument

        parser.return_value = ParsedDocument("a" * 64, "MARKDOWN", ())
        chunker.return_value = ChunkingPlan((), ())
        before_status = self.document.status
        before_paragraphs = self.document.paragraphs.count()
        before_tasks = self.document.processing_tasks.count()
        payload = preview_document_chunks(self.document)
        self.assertEqual(payload["child_count"], 0)
        self.document.refresh_from_db()
        self.assertEqual(self.document.status, before_status)
        self.assertEqual(self.document.paragraphs.count(), before_paragraphs)
        self.assertEqual(self.document.processing_tasks.count(), before_tasks)

    def test_preview_owner_and_nested_resource_permissions(self):
        owner_client = APIClient()
        owner_client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=self.owner).key}")
        other_client = APIClient()
        other_client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=self.other).key}")
        url = f"/api/knowledge-bases/{self.knowledge.id}/documents/{self.document.id}/chunk-preview/"
        with patch("api.views.preview_document_chunks", return_value={"items": []}):
            self.assertEqual(owner_client.post(url, {}, format="json").status_code, 200)
        self.assertEqual(other_client.post(url, {}, format="json").status_code, 404)
        wrong_kb = KnowledgeBase.objects.create(name="stage11-other-kb", owner=self.owner)
        wrong_url = f"/api/knowledge-bases/{wrong_kb.id}/documents/{self.document.id}/chunk-preview/"
        self.assertEqual(owner_client.post(wrong_url, {}, format="json").status_code, 404)

    def test_active_task_blocks_chunking_config_changes(self):
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=self.owner).key}")
        DocumentProcessingTask.objects.create(
            document=self.document,
            task_type=DocumentProcessingTask.TaskType.REPROCESS,
            status=DocumentProcessingTask.Status.PROCESSING,
        )
        url = (
            f"/api/knowledge-bases/{self.knowledge.id}/documents/"
            f"{self.document.id}/chunking-config/"
        )
        response = client.patch(url, {"child_target_tokens": 120}, format="json")
        self.assertEqual(response.status_code, 409)
        self.document.refresh_from_db()
        self.assertEqual(self.document.child_target_tokens, 100)

    def test_switching_back_to_legacy_still_requires_reindex(self):
        self.document.status = Document.Status.SUCCESS
        self.document.indexed_chunking_signature = chunking_signature(self.document)
        self.document.save(update_fields=["status", "indexed_chunking_signature"])
        self.assertFalse(document_needs_reprocess(self.document))
        self.document.chunk_strategy = Document.ChunkStrategy.LEGACY
        self.document.save(update_fields=["chunk_strategy"])
        self.assertTrue(document_needs_reprocess(self.document))

    def test_retrieval_expands_parent_once_and_keeps_child_reference(self):
        section = DocumentSection.objects.create(
            document=self.document,
            position=1,
            content="完整父章节：网络配置、服务名和排错步骤。",
            source_block_ids=["block-parent"],
        )
        first = Paragraph.objects.create(
            document=self.document,
            parent_section=section,
            position=1,
            content="网络配置",
            embedding=[1.0, 0.0],
            chunk_type=Paragraph.ChunkType.CHILD,
            source_block_ids=["block-network"],
        )
        Paragraph.objects.create(
            document=self.document,
            parent_section=section,
            position=2,
            content="服务名称",
            embedding=[0.9, 0.1],
            chunk_type=Paragraph.ChunkType.CHILD,
            source_block_ids=["block-service"],
        )
        self.document.status = Document.Status.SUCCESS
        self.document.paragraph_count = 2
        self.document.save(update_fields=["status", "paragraph_count"])
        result = retrieve_candidates(
            self.knowledge,
            "网络",
            embedding_function=lambda texts, knowledge: [[1.0, 0.0]],
        )
        references = result.references()
        self.assertEqual(len(references), 1)
        self.assertEqual(references[0]["paragraph_id"], first.id)
        self.assertEqual(references[0]["matched_content"], "网络配置")
        self.assertEqual(references[0]["content"], section.content)

    def test_parent_expansion_falls_back_to_child_when_budget_is_tight(self):
        self.knowledge.max_context_chars = 1000
        self.knowledge.save(update_fields=["max_context_chars"])
        section = DocumentSection.objects.create(
            document=self.document,
            position=1,
            content="父章节" * 2000,
        )
        child = Paragraph.objects.create(
            document=self.document,
            parent_section=section,
            position=1,
            content="命中的精确排错步骤",
            embedding=[1.0, 0.0],
            chunk_type=Paragraph.ChunkType.CHILD,
        )
        self.document.status = Document.Status.SUCCESS
        self.document.paragraph_count = 1
        self.document.save(update_fields=["status", "paragraph_count"])
        result = retrieve_candidates(
            self.knowledge,
            "排错",
            embedding_function=lambda texts, knowledge: [[1.0, 0.0]],
        )
        reference = result.references()[0]
        self.assertEqual(reference["paragraph_id"], child.id)
        self.assertEqual(reference["content"], child.content)
        self.assertLessEqual(result.context_chars, self.knowledge.max_context_chars)


class StageElevenEvaluationTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix="stage11-eval-")
        self.root = Path(self.temp_dir.name)
        self.digest = "a" * 64
        self.manifest = self.root / "manifest.jsonl"
        self.manifest.write_text(json.dumps({
            "source_id": "docker-networking",
            "title": "Docker networking",
            "source_url": "https://docs.docker.com/network/",
            "publisher": "Docker",
            "product": "Docker Engine",
            "product_version": "2026-09 snapshot",
            "content_type": "text/html",
            "license": "Apache-2.0 documentation repository",
            "retrieved_at": "2026-09-09T00:00:00+00:00",
            "sha256": self.digest,
            "local_path": "",
            "included": True,
        }, ensure_ascii=False) + "\n", encoding="utf-8")

    def tearDown(self):
        self.temp_dir.cleanup()

    def record(self, *, case_id="stage11-case-001", query="容器如何通过服务名通信？", split="dev", status="AUTO_DRAFT"):
        return {
            "case_id": case_id,
            "query": query,
            "category": "semantic_paraphrase",
            "split": split,
            "answerable": True,
            "gold_evidence": [{
                "source_id": "docker-networking",
                "document_sha256": self.digest,
                "source_block_ids": ["blk-network"],
                "relevance": 2,
            }],
            "answer_key_points": ["容器必须位于可互通网络"],
            "review_status": status,
            "reviewed_at": "2026-09-09T00:00:00+00:00" if status == "HUMAN_APPROVED" else "",
            "notes": "",
        }

    def write_dataset(self, rows, name="dataset.jsonl"):
        path = self.root / name
        path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")
        return path

    def test_dataset_requires_stable_evidence_and_human_status_for_formal_use(self):
        path = self.write_dataset([self.record()])
        self.assertEqual(load_stage11_cases(path, manifest_path=self.manifest)[0].review_status, "AUTO_DRAFT")
        with self.assertRaisesRegex(EvaluationDataError, "人工确认"):
            load_stage11_cases(path, manifest_path=self.manifest, require_approved=True)
        broken = self.record()
        broken["gold_evidence"][0]["source_block_ids"] = []
        with self.assertRaisesRegex(EvaluationDataError, "稳定Block ID"):
            load_stage11_cases(self.write_dataset([broken], "broken.jsonl"), manifest_path=self.manifest)

    def test_duplicate_and_near_duplicate_questions_are_rejected(self):
        rows = [
            self.record(case_id="stage11-a", query="Docker 容器怎样通过服务名通信？"),
            self.record(case_id="stage11-b", query="Docker容器怎样通过服务名通信"),
        ]
        with self.assertRaisesRegex(EvaluationDataError, "重复或高度近似"):
            load_stage11_cases(self.write_dataset(rows), manifest_path=self.manifest)

    def test_recall_counts_all_gold_and_ndcg_uses_graded_relevance(self):
        case = Stage11EvaluationCase(
            case_id="stage11-metrics",
            query="query",
            category="multi_document",
            split="test",
            answerable=True,
            gold_evidence=(
                StableEvidence("docker-networking", self.digest, ("blk-high",), 2),
                StableEvidence("docker-networking", self.digest, ("blk-low",), 1),
                StableEvidence("docker-networking", self.digest, ("blk-missing",), 2),
            ),
            answer_key_points=("key",),
            review_status="HUMAN_APPROVED",
            reviewed_at="2026-09-09T00:00:00+00:00",
            notes="",
        )
        candidates = [
            SimpleNamespace(document_source_id="other", document_sha256=self.digest, source_block_ids=["none"]),
            SimpleNamespace(document_source_id="docker-networking", document_sha256=self.digest, source_block_ids=["blk-low"]),
            SimpleNamespace(document_source_id="docker-networking", document_sha256=self.digest, source_block_ids=["blk-high"]),
        ]
        metrics = evaluate_stable_ranking(candidates, case)
        self.assertEqual(metrics["hit_at_1"], 0)
        self.assertEqual(metrics["hit_at_3"], 1)
        self.assertEqual(metrics["recall_at_5_hits"], 2)
        self.assertEqual(metrics["recall_at_5_total"], 3)
        self.assertAlmostEqual(metrics["recall_at_5"], 2 / 3)
        self.assertEqual(metrics["mrr_at_10"], 0.5)
        self.assertGreater(metrics["ndcg_at_10"], 0)
        self.assertLess(metrics["ndcg_at_10"], 1)

    def test_answer_metrics_do_not_reward_rejecting_everything(self):
        answerable = Stage11EvaluationCase(
            case_id="answerable", query="q", category="procedure", split="test", answerable=True,
            gold_evidence=(StableEvidence("docker-networking", self.digest, ("blk-network",), 2),),
            answer_key_points=("key",), review_status="HUMAN_APPROVED", reviewed_at="now", notes="",
        )
        no_answer = Stage11EvaluationCase(
            case_id="no-answer", query="q2", category="no_answer", split="test", answerable=False,
            gold_evidence=(), answer_key_points=(), review_status="HUMAN_APPROVED", reviewed_at="now", notes="",
        )
        rows = [
            evaluate_answer_record(answerable, {"citations": [], "refused": True}),
            evaluate_answer_record(no_answer, {"citations": [], "refused": True}),
        ]
        summary = aggregate_answer_rows(rows)
        self.assertEqual(summary["no_answer_accuracy"]["value"], 1)
        self.assertEqual(summary["answerable_response_rate"]["value"], 0)
        self.assertEqual(summary["faithfulness"]["judge_failures"], 2)

    def test_freeze_rejects_too_small_test_set(self):
        source = self.write_dataset([
            self.record(split="test", status="HUMAN_APPROVED")
        ], "reviewed.jsonl")
        with self.assertRaisesRegex(CommandError, "至少需要100题"):
            call_command(
                "freeze_evaluation_dataset",
                dataset=str(source),
                output=str(self.root / "frozen.jsonl"),
                manifest=str(self.manifest),
            )
        self.assertFalse((self.root / "frozen.jsonl").exists())
